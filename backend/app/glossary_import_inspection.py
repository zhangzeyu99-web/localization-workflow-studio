"""Read-only glossary source discovery; row count is guidance, not an import gate."""
from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from . import db
from .archive_batch_engine import ArchiveBatchError
from .glossary_archive_batches import DIRECT_GLOSSARY_KINDS
from .languages import SOURCE_HEADER_ALIASES, require_supported_language
from .workflow.asset_import_export import _multilingual_glossary_layout
from .workflow.table_helpers import (
    LANGUAGE_ORDER,
    XLSX_IMPORT_SUFFIXES,
    UnsupportedImportFormatError,
    _auto_language_indices,
    _column_index,
    _glossary_layout,
    _json_glossary_language_error,
    _mapping_rows_to_matrix,
    _normalized_header_indices,
    _read_csv_matrix,
    _read_json_mapping_rows,
    _select_xlsx_data_sheet,
    _sheet_headers,
    _value_at,
    _worksheet_has_source_data,
)


def _source_index(headers: list[str], request: Any) -> int:
    return _column_index(_normalized_header_indices(headers), request.source_column, list(SOURCE_HEADER_ALIASES))


def inspect_glossary_import(project_id: str, request: Any) -> dict[str, Any]:
    try:
        db.get_project(project_id)
        artifact = db.get_artifact(request.artifact_id)
    except KeyError as exc:
        raise ArchiveBatchError(404, "project_or_artifact_not_found", "项目或文件不存在。") from exc
    if artifact["project_id"] != project_id:
        raise ArchiveBatchError(404, "project_or_artifact_not_found", "项目或文件不存在。")
    if artifact.get("kind") not in DIRECT_GLOSSARY_KINDS:
        raise ArchiveBatchError(400, "candidate_scan_required", "请选择术语表文件；完整语言表可先提取术语候选。")
    path = Path(artifact["path"])
    try:
        return _inspect_file(path, request)
    except FileNotFoundError as exc:
        raise ArchiveBatchError(404, "artifact_file_missing", "已存档的来源文件不存在，请重新上传后继续。") from exc
    except (OSError, BadZipFile, InvalidFileException, SyntaxError) as exc:
        raise ArchiveBatchError(400, "invalid_glossary_file", "来源文件无法读取或已损坏，请检查文件后重新上传；现有存档记录未更改。") from exc


def _inspect_file(path: Path, request: Any) -> dict[str, Any]:
    suffix = path.suffix.lower()
    explicit_languages: set[str] = set()
    warnings: list[str] = []
    if suffix == ".csv":
        headers, raw_rows = _read_csv_matrix(path)
        sheet_name = "__csv__"
    elif suffix == ".json":
        mappings = _read_json_mapping_rows(path, ("terms", "rows", "entries"))
        requires_language = any("language" in row for row in mappings)
        invalid_language_rows = 0
        for row in mappings:
            if _json_glossary_language_error(row, requires_language=requires_language):
                invalid_language_rows += 1
                continue
            if str(row.get("language") or "").strip():
                explicit_languages.add(require_supported_language(str(row["language"])))
        if invalid_language_rows:
            warnings.append(f"有 {invalid_language_rows} 条 JSON 记录的 language 缺失或无效，将在差异预览中逐条处理；原文件保留。")
        headers, raw_rows = _mapping_rows_to_matrix(mappings)
        sheet_name = "__json__"
    elif suffix in XLSX_IMPORT_SUFFIXES:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            def has_source(worksheet: Any, *, require_data: bool = True) -> bool:
                try:
                    source_index = _source_index(_sheet_headers(worksheet), request)
                except KeyError:
                    return False
                return not require_data or _worksheet_has_source_data(worksheet, source_index)

            def is_candidate(worksheet: Any, *, multilingual: bool) -> bool:
                headers = _sheet_headers(worksheet)
                mapping = {
                    field: getattr(request, field)
                    for field in ("term_key_column", "source_column", "category_column", "note_column")
                }
                try:
                    if multilingual:
                        _, source_index, _, _, targets = _multilingual_glossary_layout(headers, **mapping)
                        if not targets:
                            return False
                    else:
                        _, _, source_index, _, _, _, _ = _glossary_layout(
                            headers, **mapping, target_column=request.target_column,
                            target_alt_column=request.target_alt_column,
                            language=require_supported_language(request.language),
                        )
                except (KeyError, ValueError):
                    return False
                return any(
                    any(value is not None and str(value).strip() for value in row)
                    for row in worksheet.iter_rows(min_row=2, values_only=True)
                )

            # Match analyze's multilingual-first selection before offering
            # source-only sheets for manual mapping or an empty-file warning.
            worksheet = None
            if request.auto_languages and not (request.target_column or request.target_alt_column):
                worksheet = _select_xlsx_data_sheet(
                    workbook, request.sheet, lambda sheet: is_candidate(sheet, multilingual=True), allow_none=True,
                )
            if worksheet is None:
                worksheet = _select_xlsx_data_sheet(
                    workbook, request.sheet, lambda sheet: is_candidate(sheet, multilingual=False), allow_none=True,
                )
            if worksheet is None:
                worksheet = _select_xlsx_data_sheet(workbook, request.sheet, has_source, allow_none=True)
            if worksheet is None:
                worksheet = _select_xlsx_data_sheet(workbook, request.sheet, lambda sheet: has_source(sheet, require_data=False))
            sheet_name = worksheet.title
            headers = _sheet_headers(worksheet)
            raw_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        finally:
            workbook.close()
    else:
        raise UnsupportedImportFormatError(suffix, (".xlsx", ".csv", ".json"))
    source_index = _source_index(headers, request)
    normalized = _normalized_header_indices(headers)
    reserved = {source_index}
    columns: dict[str, Any] = {"source": headers[source_index]}
    for field, aliases in (
        ("term_key", ["id", "key", "编号", "序号"]),
        ("category", ["category", "type", "分类", "类别", "类型"]),
        ("note", ["note", "notes", "comment", "备注"]),
    ):
        index = _column_index(normalized, getattr(request, f"{field}_column"), aliases, required=False)
        columns[field] = headers[index] if index is not None else ""
        if index is not None:
            reserved.add(index)
    detected = _auto_language_indices(headers, reserved)
    if request.target_column or request.target_alt_column:
        target_index = _column_index(normalized, request.target_column or request.target_alt_column, [])
        detected = {language: pair for language, pair in detected.items() if target_index in pair}
        explicit_languages.clear()
    languages = [language for language in LANGUAGE_ORDER if language in detected or language in explicit_languages]
    columns["languages"] = {language: {"target": headers[index], "target_alt": ""} for language, (index, _) in detected.items()}
    source_rows = sum(any(value is not None and str(value).strip() for value in row) for row in raw_rows)
    nonempty_source_rows = sum(bool(_value_at(row, source_index)) for row in raw_rows)
    if source_rows > 1000:
        warnings.append(f"检测到 {source_rows} 条原表记录，支持整表导入，无需按 1000 条拆分。请确认这是已审核的术语表；若为完整语言表，建议先提取术语候选。")
    if not languages:
        warnings.append("未从列名识别出具体语种，请手动选择语言并核对目标译文列。")
    if not nonempty_source_rows:
        warnings.append("当前工作表没有非空源文行，请检查来源或列映射。")
    return {"languages": languages, "sheet": sheet_name, "source_rows": source_rows, "columns": columns, "warnings": warnings}
