from __future__ import annotations

import os
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.db as db
from app.config import DEFAULT_SETTINGS, save_settings
from app.main import app
from conftest import reset_data_root, wait_for_background_jobs


@pytest.fixture(autouse=True)
def isolated_state(monkeypatch: pytest.MonkeyPatch):
    reset_data_root(Path(os.environ["LWS_DATA_ROOT"]))
    db.init_db()
    save_settings(DEFAULT_SETTINGS)

    def reject_network(*_args, **_kwargs):
        raise AssertionError("external HTTP is forbidden in delivery history tests")

    async def reject_async_network(*_args, **_kwargs):
        raise AssertionError("external HTTP is forbidden in delivery history tests")

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", reject_network)
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", reject_async_network)
    yield
    wait_for_background_jobs()
    save_settings(DEFAULT_SETTINGS)


def write_final(path: Path, text: str = "Start Game", language: str = "EN") -> None:
    if path.suffix == ".txt":
        path.write_text(text, encoding="utf-8")
        return
    workbook = Workbook()
    workbook.active.title = "Language"
    workbook.active.append(["ID", "CN", language])
    workbook.active.append([1, "开始游戏", text])
    workbook.save(path)
    workbook.close()


def add_run(project_id: str, path: Path, *, task_id: str = "task-one", language: str = "en", status: str = "passed", input_id: str = "", final: bool = True) -> dict:
    run = db.insert_run(project_id, "translation", language, metadata={
        "translation_task_id": task_id,
        "task_origin": "quick_task" if path.suffix == ".txt" else "translation_run",
        "input_artifact_id": input_id,
        "quality_summary": {"passed": status == "passed", "hard_errors": 0 if status == "passed" else 2},
    })
    if final:
        db.add_artifact(project_id, path.name, path, "final_text" if path.suffix == ".txt" else "qa_final_workbook", run_id=run["id"])
    return db.update_run(run["id"], status=status)


def history(client: TestClient, project_id: str) -> dict:
    response = client.get(f"/api/projects/{project_id}/delivery-history")
    assert response.status_code == 200, response.text
    return response.json()


def deliver(client: TestClient, project_id: str, run_id: str) -> dict:
    response = client.post(f"/api/projects/{project_id}/delivery-package", params={"run_id": run_id})
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize("suffix", [".txt", ".xlsx"])
def test_regeneration_lists_immutable_versions_and_frozen_qa(tmp_path: Path, suffix: str) -> None:
    project = db.insert_project("history", "QA", "")
    path = tmp_path / f"final{suffix}"
    write_final(path)
    run = add_run(project["id"], path)
    with TestClient(app) as client:
        first = deliver(client, project["id"], run["id"])
        first_url = first["files"][0]["download_url"]
        first_bytes = client.get(first_url).content
        write_final(path, "Begin Game")
        db.merge_run_metadata(run["id"], {"quality_summary": {"passed": False, "hard_errors": 2}})
        db.update_run(run["id"], status="failed")
        second = deliver(client, project["id"], run["id"])
        result = history(client, project["id"])
        assert len(result["versions"]) == 2
        old = next(item for item in result["versions"] if item["qa_snapshot"]["status"] == "passed")
        current = next(item for item in result["versions"] if item["is_current"])
        assert old["qa_snapshot"]["hard_errors"] == 0
        assert old["is_current"] is False
        assert current["qa_snapshot"]["hard_errors"] == 2
        assert all(item["task_id"] == "task-one" and item["history_complete"] for item in result["versions"])
        assert client.get(old["files"][0]["download_url"]).content == first_bytes
        assert client.get(first_url).content == first_bytes
        assert first_url != second["files"][0]["download_url"]
        assert result["current_tasks"][0]["current_version_id"] == current["version_id"]


@pytest.mark.parametrize("status", ["failed", "running", "queued", "needs_input"])
def test_new_current_run_without_output_never_uses_old_passed_delivery(tmp_path: Path, status: str) -> None:
    project = db.insert_project("current failed", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    old = add_run(project["id"], path, input_id="same-input")
    with TestClient(app) as client:
        deliver(client, project["id"], old["id"])
        # Historical databases can predate terminal-task guards; preserve their actual IDs.
        current = add_run(project["id"], path, task_id="", input_id="same-input", status=status, final=False)
        db.merge_run_metadata(current["id"], {"translation_task_id": "task-one"})
        result = history(client, project["id"])
        assert len(result["current_tasks"]) == 1
        task = result["current_tasks"][0]
        assert task["run_id"] == current["id"]
        assert task["status"] == status
        assert task["can_generate"] is False
        assert task["current_version_id"] is None
        assert result["versions"][0]["is_current"] is False
        assert result["versions"][0]["qa_snapshot"]["status"] == "passed"


def test_same_input_does_not_merge_distinct_or_unidentified_tasks(tmp_path: Path) -> None:
    project = db.insert_project("task identity", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    runs = [add_run(project["id"], path, task_id=task_id, input_id="same-input") for task_id in ["task-one", "task-two", ""]]
    with TestClient(app) as client:
        for run in runs:
            deliver(client, project["id"], run["id"])
        result = history(client, project["id"])
        assert {task["task_id"] for task in result["current_tasks"]} == {"task-one", "task-two"}
        legacy = next(version for version in result["versions"] if version["task_id"] is None)
        assert legacy["history_complete"] is False
        assert legacy["is_current"] is False


def test_legacy_generated_file_without_snapshot_is_explicitly_incomplete(tmp_path: Path) -> None:
    project = db.insert_project("legacy", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    run = add_run(project["id"], path)
    with TestClient(app) as client:
        package = deliver(client, project["id"], run["id"])
        anchor = next(item for item in db.list_artifacts(run_id=run["id"]) if item["kind"] == "delivery_version")
        db.update_artifact(anchor["id"], {"metadata": {}})
        result = history(client, project["id"])
        assert len(result["versions"]) == 1
        legacy = result["versions"][0]
        assert legacy["history_complete"] is False
        assert legacy["qa_snapshot"] == {"status": "unknown", "hard_errors": None, "soft_warnings": None}
        assert legacy["is_current"] is False
        assert client.get(legacy["files"][0]["download_url"]).content == client.get(package["files"][0]["download_url"]).content


def test_merged_deliveries_register_each_version_and_become_history_on_new_failure(tmp_path: Path) -> None:
    project = db.insert_project("merged history", "QA", "")
    source_path = tmp_path / "source.xlsx"
    write_final(source_path, "")
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    path = tmp_path / "en.xlsx"
    write_final(path)
    add_run(project["id"], path, input_id=source["id"])
    with TestClient(app) as client:
        url = f"/api/projects/{project['id']}/delivery-package/merged"
        payload = {"input_artifact_id": source["id"], "languages": ["en"], "translation_task_id": "task-one"}
        first = client.post(url, json=payload)
        second = client.post(url, json=payload)
        assert first.status_code == second.status_code == 200
        result = history(client, project["id"])
        versions = [item for item in result["versions"] if item["task_kind"] == "merged"]
        assert len(versions) == 2
        assert sum(item["is_current"] for item in versions) == 1
        first_url = first.json()["files"][0]["download_url"]
        first_bytes = client.get(first_url).content
        current = add_run(project["id"], path, task_id="", input_id=source["id"], status="failed", final=False)
        db.merge_run_metadata(current["id"], {"translation_task_id": "task-one"})
        result = history(client, project["id"])
        assert not any(item["is_current"] for item in result["versions"])
        assert next(item for item in result["current_tasks"] if item["task_kind"] == "merged")["qa_status"] == "failed"
        rejected = client.post(url, json=payload)
        assert rejected.status_code == 409, rejected.text
        assert client.get(first_url).content == first_bytes
        assert next(item for item in result["current_tasks"] if item["task_kind"] == "translation")["run_id"] == current["id"]


def test_announcement_force_regeneration_retains_package_and_qa_snapshot(tmp_path: Path) -> None:
    with TestClient(app) as client:
        project = client.post("/api/projects", json={"name": "announcement history", "type": "announcement"}).json()
        created = client.post(f"/api/projects/{project['id']}/announcement-tasks", json={"text": "开始游戏", "languages": ["en"]})
        assert created.status_code == 200, created.text
        task = created.json()
        output_path = tmp_path / "output.txt"
        write_final(output_path)
        output = db.add_artifact(project["id"], "output", output_path, "announcement_output_file", metadata={"task_id": task["id"], "language": "en"})
        for hard in [0, 2]:
            qa_path = tmp_path / f"qa-{hard}.xlsx"
            write_final(qa_path)
            qa = db.add_artifact(project["id"], "qa", qa_path, "announcement_qa_summary", metadata={"task_id": task["id"], "hard_blockers": hard})
            current = db.get_announcement_task(task["id"])
            db.update_announcement_task(task["id"], metadata={**current["metadata"], "qa_summary_artifact_id": qa["id"],
                                        "output_artifact_ids": {"en": output["id"]}, "hard_blockers": hard}, allow_terminal_update=True)
            if hard:
                write_final(output_path, "Begin Game")
            response = client.post(f"/api/announcement-tasks/{task['id']}/deliver", json={"languages": ["en"], "force": bool(hard), "date_stamp": "20260921"})
            assert response.status_code == 200, response.text
            if not hard:
                result = history(client, project["id"])
                assert len(result["versions"]) == 1
                old = result["versions"][0]
                assert old["qa_snapshot"]["status"] == "passed"
                old_bytes = client.get(old["files"][0]["download_url"]).content
        result = history(client, project["id"])
        assert len(result["versions"]) == 2
        assert sum(item["is_current"] for item in result["versions"]) == 1
        current = next(item for item in result["versions"] if item["is_current"])
        assert current["qa_snapshot"]["status"] == "failed"
        assert current["qa_snapshot"]["hard_errors"] == 2
        assert result["current_tasks"][0]["qa_status"] == "failed"
        assert result["current_tasks"][0]["task_id"] == task["id"]
        assert client.get(old["files"][0]["download_url"]).content == old_bytes
        historical = next(item for item in result["versions"] if not item["is_current"])
        assert historical["qa_snapshot"]["status"] == "passed"
        assert db.list_translation_entries(project["id"]) == []


@pytest.mark.parametrize("lineage", ["source_run_id", "manual_fix_source_run_id", "model_fix_source_run_id", ""])
def test_equal_execution_time_requires_explicit_successor_or_incomplete_current(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, lineage: str) -> None:
    monkeypatch.setattr(db, "now_iso", lambda: "2026-09-21T00:00:00+00:00")
    project = db.insert_project("tie", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    old = add_run(project["id"], path, input_id="input-a")
    current = add_run(project["id"], path, input_id="input-b", status="failed", final=False)
    if lineage:
        db.merge_run_metadata(current["id"], {lineage: old["id"]})
    with TestClient(app) as client:
        deliver(client, project["id"], old["id"])
        result = history(client, project["id"])
        task = result["current_tasks"][0]
        assert len(result["current_tasks"]) == 1
        assert task["current_version_id"] is None
        assert not any(item["is_current"] for item in result["versions"])
        if lineage:
            assert task["run_id"] == current["id"]
            assert task["status"] == "failed"
        else:
            assert task["status"] == "needs_input"
            assert task["qa_status"] == "unknown"
            assert task["current_evidence_complete"] is False
        assert task["can_generate"] is False


def test_queued_retry_supersedes_old_snapshot_but_metadata_edit_does_not(tmp_path: Path) -> None:
    project = db.insert_project("retry", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    run = add_run(project["id"], path)
    with TestClient(app) as client:
        deliver(client, project["id"], run["id"])
        first = history(client, project["id"])["current_tasks"][0]["current_version_id"]
        db.merge_run_metadata(run["id"], {"unrelated_ui_label": "new label"})
        assert history(client, project["id"])["current_tasks"][0]["current_version_id"] == first
        db.merge_run_metadata(run["id"], {"queued_at": "2099-01-01T00:00:00+00:00"})
        result = history(client, project["id"])
        assert result["current_tasks"][0]["current_version_id"] is None
        assert result["versions"][0]["is_current"] is False


def test_missing_delivery_file_has_no_download_and_cannot_be_current(tmp_path: Path) -> None:
    project = db.insert_project("missing", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    run = add_run(project["id"], path)
    with TestClient(app) as client:
        package = deliver(client, project["id"], run["id"])
        # Remove only this test's generated file inside the conftest TEMP root.
        generated = Path(package["files"][0]["path"]).resolve()
        assert generated.is_relative_to(Path(os.environ["LWS_DATA_ROOT"]).resolve())
        generated.unlink()
        result = history(client, project["id"])
        version = result["versions"][0]
        assert version["available"] is False
        assert version["files"][0]["available"] is False
        assert version["files"][0]["download_url"] == ""
        assert version["is_current"] is False
        assert result["current_tasks"][0]["current_version_id"] is None


def test_new_final_artifact_with_same_run_and_qa_invalidates_previous_delivery(tmp_path: Path) -> None:
    project = db.insert_project("new final evidence", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    run = add_run(project["id"], path)
    with TestClient(app) as client:
        deliver(client, project["id"], run["id"])
        replacement = tmp_path / "new-final.txt"
        write_final(replacement, "Begin Game")
        db.add_artifact(project["id"], "replacement", replacement, "final_text", run_id=run["id"])
        result = history(client, project["id"])
        assert result["current_tasks"][0]["current_version_id"] is None
        assert result["versions"][0]["is_current"] is False


@pytest.mark.parametrize("invalid_count", [None, "unknown", True, -1, "NaN"])
def test_invalid_current_qa_count_is_unknown_and_cannot_reuse_old_pass(tmp_path: Path, invalid_count) -> None:
    project = db.insert_project("invalid QA", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    run = add_run(project["id"], path)
    with TestClient(app) as client:
        deliver(client, project["id"], run["id"])
        db.merge_run_metadata(run["id"], {"quality_summary": {"passed": True, "hard_errors": invalid_count}})
        result = history(client, project["id"])
        task = result["current_tasks"][0]
        assert task["qa_status"] == "unknown"
        assert task["qa_hard_errors"] is None
        assert task["can_generate"] is False
        assert task["current_version_id"] is None
        assert result["versions"][0]["qa_snapshot"]["status"] == "passed"


def test_merged_partial_delivery_never_displays_as_fully_passed(tmp_path: Path) -> None:
    project = db.insert_project("partial merged", "QA", "")
    source_path = tmp_path / "source.xlsx"
    write_final(source_path, "")
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    path = tmp_path / "en.xlsx"
    write_final(path)
    add_run(project["id"], path, input_id=source["id"])
    add_run(project["id"], tmp_path / "missing-ko.xlsx", language="ko", input_id=source["id"])
    with TestClient(app) as client:
        response = client.post(f"/api/projects/{project['id']}/delivery-package/merged", json={
            "input_artifact_id": source["id"], "languages": ["en", "ko"], "translation_task_id": "task-one",
        })
        assert response.status_code == 200, response.text
        result = history(client, project["id"])
        task = next(item for item in result["current_tasks"] if item["task_kind"] == "merged")
        assert task["qa_status"] == "mixed"
        assert task["skipped_languages"] == ["KR"]
        assert result["versions"][0]["qa_snapshot"]["status"] == "mixed"
        for run in db.list_runs(project["id"]):
            db.merge_run_metadata(run["id"], {"translation_task_state": "closed"})
        assert not any(item["is_current"] for item in history(client, project["id"])["versions"])


def test_invalid_announcement_qa_is_unknown_not_zero_pass(tmp_path: Path) -> None:
    with TestClient(app) as client:
        project = client.post("/api/projects", json={"name": "invalid announcement QA", "type": "announcement"}).json()
        task = client.post(f"/api/projects/{project['id']}/announcement-tasks", json={"text": "开始游戏", "languages": ["en"]}).json()
        path = tmp_path / "qa.xlsx"
        write_final(path)
        qa = db.add_artifact(project["id"], "invalid qa", path, "announcement_qa_summary", metadata={"task_id": task["id"], "hard_blockers": "unknown"})
        db.update_announcement_task(task["id"], metadata={**task["metadata"], "qa_summary_artifact_id": qa["id"], "hard_blockers": 0})
        current = history(client, project["id"])["current_tasks"][0]
        assert current["qa_status"] == "unknown"
        assert current["qa_hard_errors"] is None


def test_legacy_unknown_generation_time_sorts_after_registered_versions(tmp_path: Path) -> None:
    project = db.insert_project("legacy sorting", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    old = add_run(project["id"], path)
    new = add_run(project["id"], path, task_id="task-two")
    with TestClient(app) as client:
        deliver(client, project["id"], old["id"])
        anchor = next(item for item in db.list_artifacts(run_id=old["id"]) if item["kind"] == "delivery_version")
        db.update_artifact(anchor["id"], {"metadata": {}})
        deliver(client, project["id"], new["id"])
        versions = history(client, project["id"])["versions"]
        assert len(versions) == 2
        assert versions[0]["generated_at"]
        assert versions[1]["generated_at"] == ""
        assert versions[1]["history_complete"] is False


@pytest.mark.parametrize("state", ["canceled", "abandoned", "closed"])
def test_closed_task_cannot_generate_or_claim_previous_delivery_current(tmp_path: Path, state: str) -> None:
    project = db.insert_project("closed task", "QA", "")
    path = tmp_path / "final.txt"
    write_final(path)
    run = add_run(project["id"], path)
    with TestClient(app) as client:
        deliver(client, project["id"], run["id"])
        db.merge_run_metadata(run["id"], {"translation_task_state": state})
        result = history(client, project["id"])
        assert result["current_tasks"][0]["task_state"] == state
        assert result["current_tasks"][0]["can_generate"] is False
        assert result["versions"][0]["is_current"] is False
