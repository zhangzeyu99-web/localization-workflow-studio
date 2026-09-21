"""Product acceptance gates must follow the committed workflow source."""
from __future__ import annotations

from pathlib import Path
import runpy

from openpyxl import Workbook

from app.workflow.large_text import cache_lint_rows, readback_gate_files


def test_cache_lint_blocks_source_bound_regressions() -> None:
    cases = [
        ("持续3小时", "Lasts 3 minutes", "quantity_duration_mismatch"),
        ("无法领取", "Can claim", "semantic_negation_missing"),
        ("点击{item}{item}", "Click {item}", "protected_token_missing"),
        ("领取奖励", "Claim reward<@1>", "protected_token_extra"),
        ("你好", "Hello！", "fullwidth_punctuation"),
        ("不能领取", "You can’t claim", "incompatible_punctuation"),
    ]
    for source, target, expected in cases:
        result = cache_lint_rows([{"id": "1", "cn": source, "translations": {"en": target}}], target_languages=["en"])
        assert expected in result["hard_by_type"], (source, result)


def test_delivery_readback_checks_source_bound_values_and_skips_styled_empty_rows(tmp_path: Path) -> None:
    path = tmp_path / "final.xlsx"
    book = Workbook()
    book.active.append(["ID", "CN", "EN"])
    book.active.append([1, "持续3小时", "Lasts 3 minutes"])
    book.active.cell(20, 1).number_format = "@"
    book.save(path)
    book.close()
    result = readback_gate_files([path], target_languages=["en"])
    assert result["hard_by_type"] == {"quantity_duration_mismatch": 1}
    assert result["checked_target_cells"] == 1
    assert not result["readback_verified"]


def test_generated_product_rules_have_no_drift() -> None:
    root = Path(__file__).resolve().parents[2]
    generator = runpy.run_path(str(root / "scripts/sync_large_text_product_rules.py"))
    assert generator["sync_product_rules"](check=True) == 0


def test_cache_lint_rejects_empty_or_unresolved_inputs() -> None:
    assert not cache_lint_rows([], target_languages=["en"])["ok_to_apply"]
    assert not cache_lint_rows([{"id": "1", "cn": "领取"}], target_languages=[])["ok_to_apply"]
    row = {"id": "1", "cn": "领取", "translations": {"EN": "Claim"}, "unresolved_issues": ["needs review"]}
    assert "unresolved_review_finding" in cache_lint_rows([row], target_languages=["en"])["hard_by_type"]


def test_delivery_readback_requires_source_and_unambiguous_target(tmp_path: Path) -> None:
    for index, headers in enumerate((["EN"], ["ID", "CN", "VN", "VI"])):
        path = tmp_path / f"invalid-{index}.xlsx"
        book = Workbook()
        book.active.append(headers)
        book.active.append(["text"] * len(headers))
        book.save(path)
        book.close()
        result = readback_gate_files([path], target_languages=["en" if index == 0 else "vn"])
        expected = "source_column_missing" if index == 0 else "ambiguous_target_column"
        assert expected in result["hard_by_type"]
        assert not result["readback_verified"]
    assert not readback_gate_files([], target_languages=["en"])["readback_verified"]


def test_vn_delivery_accepts_vi_alias_without_treating_id_as_indonesian(tmp_path: Path) -> None:
    path = tmp_path / "vietnamese.xlsx"
    book = Workbook()
    book.active.append(["ID", "CN", "VI"])
    book.active.append([1, "你好", "Xin chào"])
    book.save(path)
    book.close()
    assert readback_gate_files([path], target_languages=["VN"])["readback_verified"]
