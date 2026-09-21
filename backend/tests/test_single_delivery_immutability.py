from __future__ import annotations

import hashlib
import os
import shutil
import socket
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import app.db as db
from app.config import DEFAULT_SETTINGS, save_settings
from app.main import app
from app.workflow.common import project_dir
from conftest import reset_data_root, wait_for_background_jobs


@pytest.fixture(autouse=True)
def isolated_test_state(monkeypatch: pytest.MonkeyPatch):
    reset_data_root(Path(os.environ["LWS_DATA_ROOT"]))
    db.init_db()
    save_settings(DEFAULT_SETTINGS)
    original_connect = socket.socket.connect

    def reject_external_connect(sock, address):
        if isinstance(address, tuple) and address[0] not in {"127.0.0.1", "::1", "localhost"}:
            raise AssertionError(f"External network is forbidden in delivery tests: {address}")
        return original_connect(sock, address)

    monkeypatch.setattr(socket.socket, "connect", reject_external_connect)
    monkeypatch.setattr(db, "now_iso", lambda: "2026-09-20T10:20:15+00:00")
    yield
    wait_for_background_jobs()
    save_settings(DEFAULT_SETTINGS)


def write_workbook(path: Path, translation: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Language"
    sheet.append(["ID", "CN", "EN"])
    sheet.append([1, "开始游戏", translation])
    workbook.save(path)
    workbook.close()


def add_final_run(
    project_id: str,
    input_id: str,
    final_path: Path,
    *,
    source_run_id: str = "",
    task_origin: str = "translation_run",
    lineage_key: str = "manual_fix_source_run_id",
    hard_errors: int = 0,
) -> dict:
    metadata = {
        "input_artifact_id": input_id,
        "parent_input_artifact_id": input_id,
        "task_origin": task_origin,
        "quality_summary": {"passed": not hard_errors, "hard_errors": hard_errors},
    }
    if source_run_id:
        metadata[lineage_key] = source_run_id
    run = db.insert_run(project_id, "qa" if source_run_id else "translation", "en", metadata=metadata)
    final_kind = "final_text" if final_path.suffix == ".txt" else "qa_final_workbook"
    db.add_artifact(project_id, final_path.name, final_path, final_kind, run_id=run["id"])
    return db.update_run(run["id"], status="failed" if hard_errors else "passed")


def build_package(client: TestClient, project_id: str, run_id: str) -> dict:
    response = client.post(f"/api/projects/{project_id}/delivery-package", params={"run_id": run_id})
    assert response.status_code == 200, response.text
    return response.json()


def download_final(client: TestClient, package: dict) -> tuple[str, bytes]:
    file = next(item for item in package["files"] if item["kind"] == "final")
    response = client.get(file["download_url"])
    assert response.status_code == 200, response.text
    return file["download_url"], response.content


@pytest.mark.parametrize("task_origin", ["translation_run", "quick_task"])
@pytest.mark.parametrize("lineage_key", ["source_run_id", "manual_fix_source_run_id", "model_fix_source_run_id"])
def test_same_minute_repair_delivery_keeps_previous_download_bytes(tmp_path: Path, task_origin: str, lineage_key: str) -> None:
    project = db.insert_project("single delivery versions", "QA", "")
    source_path = tmp_path / "source.xlsx"
    first_path = tmp_path / "first.xlsx"
    repaired_path = tmp_path / "repaired.xlsx"
    write_workbook(source_path, "")
    write_workbook(first_path, "Start Game")
    write_workbook(repaired_path, "Begin Game")
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    first_run = add_final_run(project["id"], source["id"], first_path, task_origin=task_origin)
    repaired_run = add_final_run(
        project["id"], source["id"], repaired_path,
        source_run_id=first_run["id"], task_origin=task_origin, lineage_key=lineage_key,
    )
    assert first_run["created_at"] == repaired_run["created_at"]

    with TestClient(app) as client:
        first_package = build_package(client, project["id"], first_run["id"])
        first_url, first_bytes = download_final(client, first_package)
        first_hash = hashlib.sha256(first_bytes).hexdigest()

        repaired_package = build_package(client, project["id"], repaired_run["id"])
        repaired_url, repaired_bytes = download_final(client, repaired_package)
        old_download = client.get(first_url)
        assert old_download.status_code == 200
        assert hashlib.sha256(old_download.content).hexdigest() == first_hash
        assert first_url != repaired_url
        assert first_package["deliverable"]["task_id"] == repaired_package["deliverable"]["task_id"]
        assert first_package["deliverable"]["task_code"] == repaired_package["deliverable"]["task_code"] == "T"
        if task_origin == "quick_task":
            assert first_package["archive"] is None
            assert repaired_package["archive"] is None
        else:
            assert first_package["archive"]["imported_count"] == 1
            assert repaired_package["archive"] is not None
        listed = client.get(f"/api/projects/{project['id']}/deliverables")
        assert listed.status_code == 200
        by_run = {item["run_id"]: item for item in listed.json()["deliverables"]}
        assert by_run[first_run["id"]]["files"]["final"]["download_url"] == first_url
        assert by_run[repaired_run["id"]]["files"]["final"]["download_url"] == repaired_url

    for content, expected in ((first_bytes, "Start Game"), (repaired_bytes, "Begin Game")):
        workbook = load_workbook(BytesIO(content), data_only=True)
        try:
            assert workbook.active.cell(2, 3).value == expected
        finally:
            workbook.close()


def test_quick_text_repair_delivery_keeps_previous_download_bytes(tmp_path: Path) -> None:
    project = db.insert_project("quick text delivery versions", "quick-task", "")
    source_path = tmp_path / "source.txt"
    first_path = tmp_path / "first.txt"
    repaired_path = tmp_path / "repaired.txt"
    source_path.write_text("开始游戏", encoding="utf-8")
    first_path.write_text("Start Game", encoding="utf-8")
    repaired_path.write_text("Begin Game", encoding="utf-8")
    source = db.add_artifact(project["id"], "source", source_path, "quick_input")
    first_run = add_final_run(project["id"], source["id"], first_path, task_origin="quick_task")
    repaired_run = add_final_run(
        project["id"], source["id"], repaired_path,
        source_run_id=first_run["id"], task_origin="quick_task",
    )

    with TestClient(app) as client:
        first_package = build_package(client, project["id"], first_run["id"])
        first_url, first_bytes = download_final(client, first_package)
        repaired_package = build_package(client, project["id"], repaired_run["id"])
        repaired_url, repaired_bytes = download_final(client, repaired_package)
        old_download = client.get(first_url)
        assert old_download.status_code == 200
        assert hashlib.sha256(old_download.content).digest() == hashlib.sha256(first_bytes).digest()
        assert first_url != repaired_url
        assert first_bytes == b"Start Game"
        assert repaired_bytes == b"Begin Game"
        for package in (first_package, repaired_package):
            assert [item["kind"] for item in package["files"]] == ["final"]
            assert package["archive"] is None


def test_repair_delivery_keeps_previous_changes_and_qa_summary_bytes(tmp_path: Path) -> None:
    project = db.insert_project("delivery companion versions", "QA", "")
    source_path = tmp_path / "source.xlsx"
    final_path = tmp_path / "final.xlsx"
    write_workbook(source_path, "")
    write_workbook(final_path, "Start Game")
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    first_run = add_final_run(project["id"], source["id"], final_path, hard_errors=1)
    repaired_run = add_final_run(project["id"], source["id"], final_path, source_run_id=first_run["id"], hard_errors=2)
    for run, text in ((first_run, "Original changes"), (repaired_run, "Repair changes")):
        changes_path = tmp_path / f"{run['id']}-changes.xlsx"
        write_workbook(changes_path, text)
        db.add_artifact(project["id"], "changes", changes_path, "qa_changes", run_id=run["id"])

    with TestClient(app) as client:
        first_package = build_package(client, project["id"], first_run["id"])
        previous = {}
        for file in first_package["files"]:
            response = client.get(file["download_url"])
            assert response.status_code == 200
            previous[file["kind"]] = (file["download_url"], hashlib.sha256(response.content).digest())
        assert set(previous) == {"final", "changes", "qa_summary"}
        repaired_package = build_package(client, project["id"], repaired_run["id"])
        for file in repaired_package["files"]:
            old_url, old_hash = previous[file["kind"]]
            old_download = client.get(old_url)
            assert old_download.status_code == 200
            assert hashlib.sha256(old_download.content).digest() == old_hash
            assert file["download_url"] != old_url
            if file["kind"] in {"changes", "qa_summary"}:
                new_download = client.get(file["download_url"])
                assert new_download.status_code == 200
                assert hashlib.sha256(new_download.content).digest() != old_hash
        assert first_package["deliverable"]["qa_hard_errors"] == 1
        assert repaired_package["deliverable"]["qa_hard_errors"] == 2


def test_same_run_regeneration_keeps_previous_download_bytes(tmp_path: Path) -> None:
    project = db.insert_project("same run delivery versions", "QA", "")
    source_path = tmp_path / "source.xlsx"
    final_path = tmp_path / "final.xlsx"
    write_workbook(source_path, "")
    write_workbook(final_path, "Start Game")
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    run = add_final_run(project["id"], source["id"], final_path)

    with TestClient(app) as client:
        first_package = build_package(client, project["id"], run["id"])
        first_url, first_bytes = download_final(client, first_package)
        write_workbook(final_path, "Begin Game")
        regenerated_package = build_package(client, project["id"], run["id"])
        regenerated_url, regenerated_bytes = download_final(client, regenerated_package)
        old_download = client.get(first_url)
        assert old_download.status_code == 200
        assert hashlib.sha256(old_download.content).digest() == hashlib.sha256(first_bytes).digest()
        assert regenerated_url != first_url
        assert regenerated_bytes != first_bytes
        listed = client.get(f"/api/projects/{project['id']}/deliverables")
        assert listed.status_code == 200
        current = next(item for item in listed.json()["deliverables"] if item["run_id"] == run["id"])
        assert current["files"]["final"]["download_url"] == regenerated_url
        assert current["task_id"] == first_package["deliverable"]["task_id"]


def test_legacy_download_remains_readable_after_versioned_delivery(tmp_path: Path) -> None:
    project = db.insert_project("legacy delivery", "QA", "")
    source_path = tmp_path / "source.xlsx"
    final_path = tmp_path / "final.xlsx"
    write_workbook(source_path, "")
    write_workbook(final_path, "Start Game")
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    run = add_final_run(project["id"], source["id"], final_path)

    with TestClient(app) as client:
        expected = client.get(f"/api/projects/{project['id']}/deliverables").json()["deliverables"][0]
        legacy_path = project_dir(project["id"]) / "delivery" / expected["files"]["final"]["filename"]
        shutil.copy2(final_path, legacy_path)
        legacy = client.get(f"/api/projects/{project['id']}/deliverables").json()["deliverables"][0]
        legacy_url = legacy["files"]["final"]["download_url"]
        legacy_download = client.get(legacy_url)
        assert legacy_download.status_code == 200
        write_workbook(final_path, "Begin Game")
        current_package = build_package(client, project["id"], run["id"])
        current_url, _ = download_final(client, current_package)
        old_download = client.get(legacy_url)
        assert old_download.status_code == 200
        assert old_download.content == legacy_download.content
        assert current_url != legacy_url


def test_failed_regeneration_keeps_previous_delivery_pointer(tmp_path: Path) -> None:
    project = db.insert_project("failed regeneration", "QA", "")
    source_path = tmp_path / "source.xlsx"
    final_path = tmp_path / "final.xlsx"
    for path, translations in ((source_path, ("", "")), (final_path, ("Start Game", "Claim Rewards"))):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Language"
        sheet.append(["ID", "CN", "EN"])
        sheet.append([1, "开始游戏", translations[0]])
        sheet.append([2, "领取奖励", translations[1]])
        workbook.save(path)
        workbook.close()
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    run = add_final_run(project["id"], source["id"], final_path)

    with TestClient(app) as client:
        first_package = build_package(client, project["id"], run["id"])
        first_url, first_bytes = download_final(client, first_package)
        workbook = load_workbook(final_path)
        workbook.active.cell(3, 3).value = None
        workbook.save(final_path)
        workbook.close()
        failed = client.post(f"/api/projects/{project['id']}/delivery-package", params={"run_id": run["id"]})
        assert failed.status_code == 409, failed.text
        old_download = client.get(first_url)
        assert old_download.status_code == 200
        assert old_download.content == first_bytes
        listed = client.get(f"/api/projects/{project['id']}/deliverables").json()["deliverables"][0]
        assert listed["files"]["final"]["download_url"] == first_url


def test_long_project_name_delivery_downloads_every_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    if os.name == "nt":
        # 隔离 TEMP 前缀可能触发 MAX_PATH；本回归只验证文件名与下载路由的一致性。
        def long_project_dir(project_id: str) -> Path:
            return Path("\\\\?\\" + str(project_dir(project_id).resolve()))

        monkeypatch.setattr("app.workflow.project_dir", long_project_dir)
        monkeypatch.setattr("app.routers.delivery.project_dir", long_project_dir)
    project = db.insert_project("P" * 120, "QA", "")
    source_path = tmp_path / "source.xlsx"
    final_path = tmp_path / "final.xlsx"
    write_workbook(source_path, "")
    write_workbook(final_path, "Start Game")
    source = db.add_artifact(project["id"], "source", source_path, "language_table")
    run = add_final_run(project["id"], source["id"], final_path, hard_errors=1)

    try:
        with TestClient(app) as client:
            package = build_package(client, project["id"], run["id"])
            assert {file["kind"] for file in package["files"]} == {"final", "changes", "qa_summary"}
            for file in package["files"]:
                response = client.get(file["download_url"])
                assert response.status_code == 200, (file["filename"], response.text)
                assert len(file["filename"]) <= 180
                workbook = load_workbook(BytesIO(response.content), read_only=True)
                workbook.close()
            listed = client.get(f"/api/projects/{project['id']}/deliverables").json()["deliverables"][0]
            assert {file["download_url"] for file in listed["files"].values()} == {
                file["download_url"] for file in package["files"]
            }
    finally:
        if os.name == "nt":
            test_project_dir = project_dir(project["id"]).resolve()
            assert test_project_dir.is_relative_to(Path(os.environ["LWS_DATA_ROOT"]).resolve())
            shutil.rmtree(long_project_dir(project["id"]))
