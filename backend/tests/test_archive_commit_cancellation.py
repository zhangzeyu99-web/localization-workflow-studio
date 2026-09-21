from __future__ import annotations

import json
import os
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import archive_batch_engine, db, job_queue, translation_archive_batches, workflow
from app.config import DEFAULT_SETTINGS, save_settings
from app.main import app
from app.providers import TranslationItem
from conftest import reset_data_root, wait_for_background_jobs


@pytest.fixture(autouse=True)
def isolated_runtime():
    reset_data_root(Path(os.environ["LWS_DATA_ROOT"]))
    db.init_db()
    save_settings(DEFAULT_SETTINGS)
    yield
    wait_for_background_jobs()
    save_settings(DEFAULT_SETTINGS)


@pytest.mark.parametrize("kind", ["qa", "translation"])
@pytest.mark.parametrize("boundary", ["after_analysis", "inside_transaction", "after_commit", "after_commit_event"])
def test_cancel_and_archive_have_one_transactional_winner(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, boundary: str, kind: str,
) -> None:
    path = tmp_path / "qa.xlsx"
    workbook = Workbook()
    workbook.active.append(["ID", "CN", "EN"])
    workbook.active.append(["start", "开始游戏", "Start Game" if kind == "qa" else ""])
    workbook.save(path)
    workbook.close()
    reached = threading.Event()
    release = threading.Event()
    worker_events: list[threading.Event] = []
    cancel_entered = threading.Event()
    module = archive_batch_engine if boundary == "inside_transaction" else translation_archive_batches
    boundary_function = {
        "after_analysis": "analyze_translation_archive", "inside_transaction": "_reserve_state_version",
        "after_commit": "commit_translation_archive", "after_commit_event": "commit_translation_archive",
    }[boundary]
    original = getattr(module, boundary_function)
    original_cancel = job_queue.cancel_job

    def observed_cancel(*args: Any, **kwargs: Any):
        cancel_entered.set()
        return original_cancel(*args, **kwargs)

    def analyze_then_pause(*args: Any, **kwargs: Any):
        result = original(*args, **kwargs)
        if kwargs.get("cancel_event") is not None:
            worker_events.append(kwargs["cancel_event"])
        reached.set()
        assert release.wait(15), "test did not release archive analysis"
        return result

    async def translate(batch: list[dict[str, Any]], *_args: Any):
        return [TranslationItem(id=row["id"], translation="Start Game") for row in batch]

    monkeypatch.setattr(module, boundary_function, analyze_then_pause)
    monkeypatch.setattr(job_queue, "cancel_job", observed_cancel)
    monkeypatch.setattr(workflow, "translate_batch", translate)
    monkeypatch.setattr(workflow, "call_text", lambda *_args, **_kwargs: json.dumps({"passed": True, "issues": []}))
    save_settings({**DEFAULT_SETTINGS, "provider": "openai", "api_key": "test-only", "model": "local-provider-stub", "max_batch_attempts": 1})
    with TestClient(app) as client:
        project = client.post("/api/projects", json={"name": "Archive cancellation", "type": "QA"}).json()
        with path.open("rb") as file:
            artifact = client.post(
                f"/api/projects/{project['id']}/files?kind=language_table",
                files={"file": (path.name, file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            ).json()
        run = client.post("/api/runs", json={
            "project_id": project["id"], "kind": kind, "language": "en", "input_artifact_id": artifact["id"],
        }).json()
        action = "qa" if kind == "qa" else "translate"
        response = client.post(f"/api/runs/{run['id']}/{action}/start", json={"confirm_api_budget": True, "confirm_term_gap": True})
        assert response.status_code == 200, response.text
        try:
            assert reached.wait(20), client.get(f"/api/runs/{run['id']}").json()
            if boundary == "after_commit_event":
                worker_events[0].set()
                canceled = None
            elif boundary == "inside_transaction":
                with ThreadPoolExecutor(max_workers=1) as pool:
                    pending_cancel = pool.submit(client.post, f"/api/runs/{run['id']}/{action}/cancel")
                    try:
                        assert cancel_entered.wait(5)
                        with pytest.raises(TimeoutError):
                            pending_cancel.result(timeout=0.05)
                    finally:
                        release.set()
                    canceled = pending_cancel.result(timeout=15)
            else:
                if boundary == "after_commit":
                    lanes = client.get("/api/system/job-queues").json()["lanes"]
                    queued_view = next(lane["running"] for lane in lanes if lane.get("running"))
                    assert queued_view["archive_committed"] is True
                    assert queued_view["can_cancel"] is False
                    system_cancel = client.post(f"/api/system/job-queues/{queued_view['job_id']}/cancel")
                    assert system_cancel.status_code == 409, system_cancel.text
                canceled = client.post(f"/api/runs/{run['id']}/{action}/cancel")
            if canceled is not None:
                assert canceled.status_code == (200 if boundary == "after_analysis" else 409), canceled.text
            if canceled is not None and boundary != "after_analysis":
                assert canceled.json()["detail"]["code"] == "archive_already_committed"
        finally:
            release.set()
        wait_for_background_jobs()
        entries = client.get(f"/api/projects/{project['id']}/translations").json()
        final_run = client.get(f"/api/runs/{run['id']}").json()
        final_job = job_queue.get_job(f"{'qa' if kind == 'qa' else 'run'}:{run['id']}")
        with db.connect() as conn:
            revisions = conn.execute("SELECT COUNT(*) FROM archive_import_revisions WHERE project_id = ?", (project["id"],)).fetchone()[0]
            batches = conn.execute("SELECT status FROM archive_import_batches WHERE project_id = ?", (project["id"],)).fetchall()
            version = conn.execute("SELECT version FROM archive_state_versions WHERE project_id = ? AND kind = 'translations'", (project["id"],)).fetchone()
        committed = boundary != "after_analysis"
        assert {"entries": len(entries), "revisions": revisions, "run": final_run["status"], "queue": final_job["status"]} == {
            "entries": int(committed), "revisions": int(committed),
            "run": "passed" if committed else "canceled", "queue": "completed" if committed else "canceled",
        }
        assert [row["status"] for row in batches] == ["committed" if committed else "analyzed"]
        expected_version = final_run["metadata"]["translation_archive"]["state_version"] if committed else 0
        assert int(version["version"] if version else 0) == expected_version
        assert bool(final_job["payload"].get("archive_commit")) is committed
        if kind == "qa" and boundary == "after_commit":
            model_fix = job_queue.enqueue_job(
                job_id=f"model-fix:{run['id']}", lane="language_table", job_kind="model_fix",
                project_id=project["id"], target_id=run["id"], autostart=False,
            )
            newer_cancel = client.post(f"/api/runs/{run['id']}/qa/cancel")
            assert newer_cancel.status_code == 200, newer_cancel.text
            assert job_queue.get_job(model_fix["job_id"])["status"] == "canceled"
            retried = job_queue.enqueue_job(
                job_id=final_job["job_id"], lane="language_table", job_kind="qa",
                project_id=project["id"], target_id=run["id"], autostart=False,
            )
            assert not retried["payload"].get("archive_commit")
            assert client.post(f"/api/runs/{run['id']}/qa/cancel").status_code == 200


@pytest.mark.parametrize("kind", ["qa", "translation"])
@pytest.mark.parametrize("mode", ["quick_task", "commit_failure"])
def test_no_archive_marker_for_quick_tasks_or_rolled_back_transaction(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str, mode: str,
) -> None:
    path = tmp_path / "qa.xlsx"
    workbook = Workbook()
    workbook.active.append(["ID", "CN", "EN"])
    workbook.active.append(["start", "开始游戏", "Start Game" if kind == "qa" else ""])
    workbook.save(path)
    workbook.close()
    original_dump = archive_batch_engine.json_dump

    def fail_before_commit(value: Any):
        if isinstance(value, dict) and value.get("archive_commit"):
            raise RuntimeError("injected archive commit failure")
        return original_dump(value)

    def forbid_archive(*_args: Any, **_kwargs: Any):
        raise AssertionError("quick tasks must not write the translation archive")

    async def translate(batch: list[dict[str, Any]], *_args: Any):
        return [TranslationItem(id=row["id"], translation="Start Game") for row in batch]

    if mode == "quick_task":
        monkeypatch.setattr(translation_archive_batches, "analyze_translation_archive", forbid_archive)
    else:
        monkeypatch.setattr(archive_batch_engine, "json_dump", fail_before_commit)
    monkeypatch.setattr(workflow, "translate_batch", translate)
    monkeypatch.setattr(workflow, "call_text", lambda *_args, **_kwargs: json.dumps({"passed": True, "issues": []}))
    save_settings({**DEFAULT_SETTINGS, "provider": "openai", "api_key": "test-only", "model": "local-provider-stub", "max_batch_attempts": 1})
    with TestClient(app) as client:
        project = client.post("/api/projects", json={"name": "Archive safety", "type": "QA"}).json()
        with path.open("rb") as file:
            artifact = client.post(
                f"/api/projects/{project['id']}/files?kind=language_table",
                files={"file": (path.name, file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            ).json()
        request = {"project_id": project["id"], "kind": kind, "language": "en", "input_artifact_id": artifact["id"]}
        if mode == "quick_task":
            request.update(task_origin="quick_task", translation_task_id="quick-task-archive-safety")
        run = client.post("/api/runs", json=request).json()
        action = "qa" if kind == "qa" else "translate"
        response = client.post(f"/api/runs/{run['id']}/{action}/start", json={"confirm_api_budget": True, "confirm_term_gap": True})
        assert response.status_code == 200, response.text
        wait_for_background_jobs()
        final = client.get(f"/api/runs/{run['id']}").json()
        queue = job_queue.get_job(f"{'qa' if kind == 'qa' else 'run'}:{run['id']}")
        assert final["status"] == ("passed" if mode == "quick_task" else "failed")
        assert queue["status"] == ("completed" if mode == "quick_task" else "failed")
        assert not queue["payload"].get("archive_commit")
        assert not final["metadata"].get("translation_archive")
        assert client.get(f"/api/projects/{project['id']}/translations").json() == []
        with db.connect() as conn:
            assert conn.execute("SELECT COUNT(*) FROM archive_import_revisions WHERE project_id = ?", (project["id"],)).fetchone()[0] == 0
            states = conn.execute("SELECT version FROM archive_state_versions WHERE project_id = ?", (project["id"],)).fetchall()
            assert not any(row["version"] for row in states)
            batches = conn.execute("SELECT status FROM archive_import_batches WHERE project_id = ?", (project["id"],)).fetchall()
            assert [row["status"] for row in batches] == ([] if mode == "quick_task" else ["analyzed"])


def test_multilingual_cancel_preserves_committed_child_without_sealing_parent(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    workbook.active.append(["ID", "CN", "EN", "FR"])
    workbook.active.append(["start", "开始游戏", "", ""])
    workbook.save(path)
    workbook.close()
    reached = threading.Event()
    release = threading.Event()
    original_commit = translation_archive_batches.commit_translation_archive

    def commit_then_pause(*args: Any, **kwargs: Any):
        result = original_commit(*args, **kwargs)
        reached.set()
        assert release.wait(15)
        return result

    async def translate(batch: list[dict[str, Any]], *_args: Any):
        return [TranslationItem(id=row["id"], translation="Start Game") for row in batch]

    monkeypatch.setattr(translation_archive_batches, "commit_translation_archive", commit_then_pause)
    monkeypatch.setattr(workflow, "translate_batch", translate)
    monkeypatch.setattr(workflow, "call_text", lambda *_args, **_kwargs: json.dumps({"passed": True, "issues": []}))
    save_settings({**DEFAULT_SETTINGS, "provider": "openai", "api_key": "test-only", "model": "local-provider-stub", "max_batch_attempts": 1})
    with TestClient(app) as client:
        project = client.post("/api/projects", json={"name": "Multilingual archive cancellation", "type": "QA"}).json()
        with path.open("rb") as file:
            artifact = client.post(
                f"/api/projects/{project['id']}/files?kind=language_table",
                files={"file": (path.name, file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            ).json()
        started = client.post(f"/api/projects/{project['id']}/multilingual/translate/start", json={
            "input_artifact_id": artifact["id"], "languages": ["en", "fr"], "confirm_api_budget": True, "confirm_term_gap": True,
        })
        assert started.status_code == 200, started.text
        try:
            assert reached.wait(20), started.text
            queue = job_queue.list_active_jobs(project_id=project["id"])[0]
            assert not queue["payload"].get("archive_commit")
            canceled = client.post(f"/api/system/job-queues/{queue['job_id']}/cancel")
            assert canceled.status_code == 200, canceled.text
        finally:
            release.set()
        wait_for_background_jobs()
        runs = {run["language"]: run for run in db.list_runs(project["id"])}
        assert {language: run["status"] for language, run in runs.items()} == {"en": "passed", "fr": "canceled"}
        assert job_queue.get_job(queue["job_id"])["status"] == "canceled"
        assert len(client.get(f"/api/projects/{project['id']}/translations").json()) == 1
        late_cancel = client.post(f"/api/runs/{runs['en']['id']}/translate/cancel")
        assert late_cancel.status_code == 409, late_cancel.text
        assert client.get(f"/api/runs/{runs['en']['id']}").json()["status"] == "passed"


@pytest.mark.parametrize("boundary", ["after_analysis", "after_analysis_db_flag", "after_commit", "after_commit_event"])
def test_model_fix_parent_and_qa_child_share_archive_cancellation_boundary(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, boundary: str,
) -> None:
    path = tmp_path / "fix.xlsx"
    workbook = Workbook()
    workbook.active.title = "Language"
    workbook.active.append(["ID", "CN", "EN"])
    workbook.active.append(["reward", "奖励", "Forbidden Brand"])
    workbook.save(path)
    workbook.close()
    reached = threading.Event()
    release = threading.Event()
    worker_events: list[threading.Event] = []
    name = "analyze_translation_archive" if boundary.startswith("after_analysis") else "commit_translation_archive"
    original = getattr(translation_archive_batches, name)

    def pause(*args: Any, **kwargs: Any):
        result = original(*args, **kwargs)
        if kwargs.get("cancel_event") is not None:
            worker_events.append(kwargs["cancel_event"])
        reached.set()
        assert release.wait(15)
        return result

    def provider(_settings: dict, prompt: str) -> str:
        if "待修复行" not in prompt:
            return json.dumps({"passed": True, "issues": []})
        return json.dumps({"fixes": [{
            "issue_id": "project_harness:0:Language:2:forbidden_translation", "sheet": "Language", "row": 2,
            "translation": "Reward", "note": "remove forbidden phrase",
        }]})

    monkeypatch.setattr(translation_archive_batches, name, pause)
    monkeypatch.setattr(workflow, "_call_semantic_provider", provider)
    save_settings({**DEFAULT_SETTINGS, "provider": "openai", "api_key": "test-only", "model": "local-provider-stub"})
    with TestClient(app) as client:
        project = client.post("/api/projects", json={"name": "Model fix archive cancellation", "type": "QA"}).json()
        client.patch(f"/api/projects/{project['id']}/harness", json={"forbidden_translations": ["Forbidden Brand"]})
        with path.open("rb") as file:
            artifact = client.post(
                f"/api/projects/{project['id']}/files?kind=final_workbook",
                files={"file": (path.name, file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            ).json()
        run = client.post("/api/runs", json={
            "project_id": project["id"], "kind": "qa", "language": "en", "input_artifact_id": artifact["id"],
        }).json()
        assert client.post(f"/api/runs/{run['id']}/qa").json()["run"]["status"] == "failed"
        started = client.post(f"/api/runs/{run['id']}/model-fixes/start", json={"rerun_qa": True})
        assert started.status_code == 200, started.text
        try:
            assert reached.wait(20), client.get(f"/api/runs/{run['id']}").json()
            if boundary == "after_commit_event":
                worker_events[0].set()
            elif boundary == "after_analysis_db_flag":
                event_set_entered = threading.Event()
                allow_event_set = threading.Event()
                worker_event = job_queue._RUNNING["language_table"].cancel_event
                original_set = worker_event.set
                parent_finalized = threading.Event()
                original_update = db.update_run_if_task_open

                def delayed_event_set():
                    event_set_entered.set()
                    assert allow_event_set.wait(15)
                    original_set()

                def observed_terminal(run_id: str, *args: Any, **kwargs: Any):
                    result = original_update(run_id, *args, **kwargs)
                    if run_id == run["id"] and kwargs.get("status") in {"canceled", "failed"}:
                        parent_finalized.set()
                    return result

                monkeypatch.setattr(worker_event, "set", delayed_event_set)
                monkeypatch.setattr(db, "update_run_if_task_open", observed_terminal)
                with ThreadPoolExecutor(max_workers=1) as pool:
                    pending = pool.submit(client.post, f"/api/runs/{run['id']}/qa/cancel")
                    try:
                        assert event_set_entered.wait(5)
                        assert job_queue.get_job(f"model-fix:{run['id']}")["cancel_requested"] is True
                        assert not worker_event.is_set()
                        release.set()
                        assert parent_finalized.wait(5)
                    finally:
                        allow_event_set.set()
                    assert pending.result(timeout=15).status_code == 200
            else:
                canceled = client.post(f"/api/runs/{run['id']}/qa/cancel")
                assert canceled.status_code == (200 if boundary == "after_analysis" else 409), canceled.text
        finally:
            release.set()
        wait_for_background_jobs()
        final_runs = db.list_runs(project["id"])
        queue = job_queue.get_job(f"model-fix:{run['id']}")
        committed = not boundary.startswith("after_analysis")
        assert len(final_runs) == 2
        assert {item["status"] for item in final_runs} == {"passed" if committed else "canceled"}
        assert db.get_run(run["id"])["metadata"]["model_fix_status"] == ("passed" if committed else "canceled")
        assert queue["status"] == ("completed" if committed else "canceled")
        assert bool(queue["payload"].get("archive_commit")) is committed
        assert len(client.get(f"/api/projects/{project['id']}/translations").json()) == int(committed)
        with db.connect() as conn:
            revisions = conn.execute("SELECT COUNT(*) FROM archive_import_revisions WHERE project_id = ?", (project["id"],)).fetchone()[0]
            assert revisions == int(committed)


def test_old_qa_child_recheck_cannot_seal_new_model_fix_execution(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "retry.xlsx"
    workbook = Workbook()
    workbook.active.title = "Language"
    workbook.active.append(["ID", "CN", "EN"])
    workbook.active.append(["reward", "奖励", "Forbidden Brand"])
    workbook.save(path)
    workbook.close()
    second_fix_reached = threading.Event()
    release_second_fix = threading.Event()
    first_analysis_reached = threading.Event()
    release_first_analysis = threading.Event()
    fix_calls = 0
    analysis_calls = 0
    original_analysis = translation_archive_batches.analyze_translation_archive

    def first_analysis_pause(*args: Any, **kwargs: Any):
        nonlocal analysis_calls
        result = original_analysis(*args, **kwargs)
        analysis_calls += 1
        if analysis_calls == 1:
            first_analysis_reached.set()
            assert release_first_analysis.wait(15)
        return result

    def provider(_settings: dict, prompt: str) -> str:
        nonlocal fix_calls
        if "待修复行" not in prompt:
            return json.dumps({"passed": True, "issues": []})
        fix_calls += 1
        if fix_calls == 2:
            second_fix_reached.set()
            assert release_second_fix.wait(15)
        return json.dumps({"fixes": [{
            "issue_id": "project_harness:0:Language:2:forbidden_translation", "sheet": "Language", "row": 2,
            "translation": "Reward", "note": "remove forbidden phrase",
        }]})

    monkeypatch.setattr(workflow, "_call_semantic_provider", provider)
    monkeypatch.setattr(translation_archive_batches, "analyze_translation_archive", first_analysis_pause)
    save_settings({**DEFAULT_SETTINGS, "provider": "openai", "api_key": "test-only", "model": "local-provider-stub"})
    with TestClient(app) as client:
        project = client.post("/api/projects", json={"name": "Model fix execution identity", "type": "QA"}).json()
        client.patch(f"/api/projects/{project['id']}/harness", json={"forbidden_translations": ["Forbidden Brand"]})
        with path.open("rb") as file:
            artifact = client.post(
                f"/api/projects/{project['id']}/files?kind=final_workbook",
                files={"file": (path.name, file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            ).json()
        run = client.post("/api/runs", json={
            "project_id": project["id"], "kind": "qa", "language": "en", "input_artifact_id": artifact["id"],
        }).json()
        assert client.post(f"/api/runs/{run['id']}/qa").json()["run"]["status"] == "failed"
        assert client.post(f"/api/runs/{run['id']}/model-fixes/start", json={"rerun_qa": True}).status_code == 200
        try:
            assert first_analysis_reached.wait(15)
            assert client.post(f"/api/runs/{run['id']}/qa/cancel").status_code == 200
        finally:
            release_first_analysis.set()
        wait_for_background_jobs()
        child_id = next(item["id"] for item in db.list_runs(project["id"]) if item["metadata"].get("model_fix_source_run_id") == run["id"])
        first_job = job_queue.get_job(f"model-fix:{run['id']}")
        assert first_job["status"] == "canceled"
        assert not first_job["payload"].get("archive_commit")
        assert client.post(f"/api/runs/{run['id']}/model-fixes/start", json={"rerun_qa": True}).status_code == 200
        try:
            assert second_fix_reached.wait(15)
            current_job = job_queue.get_job(first_job["job_id"])
            assert current_job["queued_at"] != first_job["queued_at"]
            assert not current_job["payload"].get("archive_commit")
            rechecked = client.post(f"/api/runs/{child_id}/qa")
            assert rechecked.status_code == 200, rechecked.text
            assert rechecked.json()["run"]["status"] == "passed"
            assert rechecked.json()["run"]["metadata"]["translation_archive"]["status"] == "committed", rechecked.text
            current_job = job_queue.get_job(first_job["job_id"])
            assert not current_job["payload"].get("archive_commit"), "old QA must not claim the current repair execution"
            canceled = client.post(f"/api/runs/{run['id']}/qa/cancel")
            assert canceled.status_code == 200, canceled.text
        finally:
            release_second_fix.set()
        wait_for_background_jobs()
        assert db.get_run(child_id)["status"] == "passed"
        assert db.get_run(run["id"])["status"] == "canceled"
        assert job_queue.get_job(first_job["job_id"])["status"] == "canceled"
